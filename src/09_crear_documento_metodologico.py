"""Etapa 10: crea el informe metodológico y de resultados en formato Word editable."""  # Describe el producto generado por este programa.

from pathlib import Path  # Gestiona rutas de entrada y salida de forma portable.
from docx import Document  # Construye el archivo Word editable.
from docx.enum.section import WD_SECTION  # Controla los saltos de sección del documento.
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # Alinea verticalmente el contenido de las tablas.
from docx.enum.text import WD_ALIGN_PARAGRAPH  # Controla la alineación de párrafos.
from docx.oxml import OxmlElement  # Permite añadir campos y propiedades OOXML no expuestas por python-docx.
from docx.oxml.ns import qn  # Convierte nombres XML al espacio de nombres de Word.
from docx.shared import Inches, Pt, RGBColor  # Define medidas, tamaños y colores.

RAIZ = Path(__file__).resolve().parents[1]  # Localiza la raíz del proyecto ENEMDU.
SALIDA = RAIZ / "outputs" / "final" / "Informe_metodologico_ENEMDU_2025_APA7.docx"  # Fija la ruta del entregable.
FIGURAS = RAIZ / "reports" / "figures"  # Localiza las figuras científicas ya validadas.
ACTIVOS = RAIZ / "assets" / "institucional"  # Localiza los activos institucionales verificados.
SALIDA.parent.mkdir(parents=True, exist_ok=True)  # Crea la carpeta final si todavía no existe.

AZUL = "17365D"  # Define el azul institucional usado solo en tablas.
AZUL_CLARO = "EAF1F8"  # Define el fondo alterno de las tablas.
GRIS = "D9D9D9"  # Define el color de bordes de las tablas.
NEGRO = RGBColor(0, 0, 0)  # Define el color negro exigido para títulos y encabezados.


def sombrear(celda, color):  # Aplica un fondo uniforme a una celda.
    propiedades = celda._tc.get_or_add_tcPr()  # Obtiene las propiedades XML de la celda.
    relleno = OxmlElement("w:shd")  # Crea la instrucción XML de sombreado.
    relleno.set(qn("w:fill"), color)  # Asigna el color hexadecimal solicitado.
    propiedades.append(relleno)  # Inserta el sombreado en la celda.


def bordes_tabla(tabla):  # Establece bordes visibles y discretos en toda la tabla.
    propiedades = tabla._tbl.tblPr  # Recupera las propiedades XML de la tabla.
    bordes = propiedades.first_child_found_in("w:tblBorders")  # Busca una definición previa de bordes.
    if bordes is None:  # Comprueba si la tabla todavía no tiene bordes explícitos.
        bordes = OxmlElement("w:tblBorders")  # Crea el contenedor XML de bordes.
        propiedades.append(bordes)  # Añade el contenedor a las propiedades de la tabla.
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):  # Recorre los seis bordes necesarios.
        borde = OxmlElement(f"w:{lado}")  # Crea el borde correspondiente.
        borde.set(qn("w:val"), "single")  # Usa una línea continua.
        borde.set(qn("w:sz"), "4")  # Fija un grosor visual ligero.
        borde.set(qn("w:color"), GRIS)  # Aplica gris claro para no recargar la tabla.
        bordes.append(borde)  # Incorpora el borde a la definición.


def agregar_numero_pagina(parrafo):  # Inserta el número automático de página en el encabezado.
    parrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT  # Alinea el número en la esquina superior derecha.
    inicio = OxmlElement("w:fldChar")  # Crea el inicio del campo automático.
    inicio.set(qn("w:fldCharType"), "begin")  # Declara el comienzo del campo.
    instruccion = OxmlElement("w:instrText")  # Crea la instrucción del campo.
    instruccion.set(qn("xml:space"), "preserve")  # Conserva los espacios de la instrucción.
    instruccion.text = "PAGE"  # Solicita a Word el número de página actual.
    separador = OxmlElement("w:fldChar")  # Crea el separador interno del campo.
    separador.set(qn("w:fldCharType"), "separate")  # Marca el cambio entre instrucción y resultado.
    resultado = OxmlElement("w:t")  # Crea un resultado provisional visible.
    resultado.text = "1"  # Muestra uno hasta que Word actualice el campo.
    cierre = OxmlElement("w:fldChar")  # Crea el cierre del campo automático.
    cierre.set(qn("w:fldCharType"), "end")  # Declara el final del campo.
    corrida = parrafo.add_run()  # Crea la corrida que contendrá el campo.
    corrida._r.extend([inicio, instruccion, separador, resultado, cierre])  # Inserta todas las piezas del campo.


def parrafo(texto="", negrita=False, cursiva=False, centrado=False, sangria=True):  # Añade prosa con el formato APA del cuerpo.
    p = documento.add_paragraph()  # Crea un nuevo párrafo.
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.JUSTIFY  # Selecciona alineación según el contenido.
    p.paragraph_format.first_line_indent = Inches(0.5) if sangria and not centrado else None  # Aplica sangría de primera línea.
    p.paragraph_format.line_spacing = 2  # Mantiene doble espacio en el texto principal.
    p.paragraph_format.space_after = Pt(0)  # Evita espacios extra incompatibles con APA.
    r = p.add_run(texto)  # Inserta el contenido textual.
    r.bold = negrita  # Aplica negrita cuando se solicita.
    r.italic = cursiva  # Aplica cursiva cuando se solicita.
    r.font.name = "Times New Roman"  # Usa una tipografía aceptada por APA 7.
    r.font.size = Pt(12)  # Fija el tamaño estándar del cuerpo.
    r.font.color.rgb = NEGRO  # Mantiene el texto en negro.
    return p  # Devuelve el párrafo para ajustes adicionales.


def encabezado(texto, nivel=1):  # Añade un encabezado con jerarquía APA.
    p = documento.add_paragraph(style=f"Heading {nivel}")  # Usa el estilo semántico de Word.
    p.paragraph_format.keep_with_next = True  # Mantiene el encabezado junto al texto siguiente.
    p.paragraph_format.space_before = Pt(12)  # Separa visualmente las secciones.
    p.paragraph_format.space_after = Pt(6)  # Mantiene proximidad con el contenido.
    r = p.add_run(texto)  # Inserta el título de la sección.
    r.font.name = "Times New Roman"  # Usa la misma familia tipográfica del informe.
    r.font.color.rgb = NEGRO  # Conserva el encabezado en negro.
    r.bold = True  # Destaca la jerarquía con negrita.
    r.font.size = Pt(14 if nivel == 1 else 12)  # Diferencia niveles sin exagerar el tamaño.
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if nivel == 1 else WD_ALIGN_PARAGRAPH.LEFT  # Aplica alineación APA por nivel.
    return p  # Devuelve el encabezado creado.


def tabla(titulo, columnas, filas, anchos=None, nota=None, mantener_junta=False):  # Construye una tabla científica editable y permite conservar juntas las tablas compactas.
    p_numero = documento.add_paragraph()  # Crea el rótulo de tabla.
    p_numero.paragraph_format.keep_with_next = True  # Evita separar el número del título.
    r_numero = p_numero.add_run(titulo[0])  # Inserta el número de tabla.
    r_numero.bold = True  # Sigue la convención APA para el número.
    r_numero.font.name = "Times New Roman"  # Mantiene la tipografía del documento.
    r_numero.font.size = Pt(11)  # Usa un tamaño compacto y legible.
    p_titulo = documento.add_paragraph()  # Crea el título descriptivo de la tabla.
    p_titulo.paragraph_format.keep_with_next = True  # Mantiene el título junto a la tabla.
    r_titulo = p_titulo.add_run(titulo[1])  # Inserta el título.
    r_titulo.italic = True  # Aplica cursiva según APA.
    r_titulo.font.name = "Times New Roman"  # Mantiene la tipografía.
    r_titulo.font.size = Pt(11)  # Conserva un tamaño legible.
    t = documento.add_table(rows=1, cols=len(columnas))  # Crea la tabla y su fila de encabezados.
    t.autofit = False  # Permite controlar proporciones de columnas.
    bordes_tabla(t)  # Aplica bordes consistentes.
    for indice, etiqueta in enumerate(columnas):  # Recorre los encabezados.
        celda = t.rows[0].cells[indice]  # Selecciona la celda correspondiente.
        celda.text = str(etiqueta)  # Inserta la etiqueta de columna.
        sombrear(celda, AZUL)  # Aplica el fondo azul oscuro.
        celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER  # Centra verticalmente el contenido.
        for run in celda.paragraphs[0].runs:  # Recorre las corridas del encabezado.
            run.font.name = "Times New Roman"  # Define la familia tipográfica.
            run.font.size = Pt(9)  # Usa un tamaño compacto para tablas.
            run.bold = True  # Resalta los encabezados.
            run.font.color.rgb = RGBColor(255, 255, 255)  # Usa texto blanco sobre azul.
        if anchos:  # Comprueba si se definieron anchos específicos.
            celda.width = Inches(anchos[indice])  # Asigna el ancho del encabezado.
    for indice_fila, valores in enumerate(filas):  # Recorre los registros de la tabla.
        celdas = t.add_row().cells  # Añade una fila editable.
        for indice_columna, valor in enumerate(valores):  # Recorre los valores de la fila.
            celdas[indice_columna].text = str(valor)  # Inserta cada valor.
            celdas[indice_columna].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER  # Centra verticalmente el contenido.
            if indice_fila % 2 == 1:  # Alterna el fondo para facilitar la lectura.
                sombrear(celdas[indice_columna], AZUL_CLARO)  # Aplica azul muy claro a filas pares visuales.
            if anchos:  # Comprueba si se definieron anchos específicos.
                celdas[indice_columna].width = Inches(anchos[indice_columna])  # Asigna el ancho previsto.
            for p in celdas[indice_columna].paragraphs:  # Recorre los párrafos internos.
                p.paragraph_format.space_after = Pt(0)  # Elimina espacios innecesarios dentro de celdas.
                p.paragraph_format.line_spacing = 1.05  # Conserva respiración sin duplicar la altura.
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if indice_columna == 0 else WD_ALIGN_PARAGRAPH.CENTER  # Alinea texto y cifras de forma intencional.
                for run in p.runs:  # Recorre el texto de la celda.
                    run.font.name = "Times New Roman"  # Define la tipografía.
                    run.font.size = Pt(9)  # Mantiene legibilidad en tablas densas.
    if mantener_junta:  # Comprueba si la tabla es lo bastante compacta para no dividirse entre páginas.
        for fila in t.rows[:-1]:  # Recorre todas las filas salvo la última para encadenarlas verticalmente.
            for celda in fila.cells:  # Recorre cada celda de la fila que debe permanecer con la siguiente.
                for p in celda.paragraphs:  # Recorre los párrafos internos que controlan el salto de página.
                    propiedades = p._p.get_or_add_pPr()  # Obtiene las propiedades XML del párrafo.
                    if propiedades.find(qn("w:keepNext")) is None:  # Evita duplicar la propiedad si ya existe.
                        propiedades.append(OxmlElement("w:keepNext"))  # Mantiene el párrafo con el contenido siguiente.
    documento.add_paragraph()  # Añade separación después de la tabla.
    if nota:  # Comprueba si la tabla necesita una nota metodológica.
        p_nota = documento.add_paragraph()  # Crea el párrafo de nota.
        p_nota.paragraph_format.line_spacing = 1  # Usa espacio sencillo en notas de tabla.
        r_etiqueta = p_nota.add_run("Nota. ")  # Inserta el prefijo APA.
        r_etiqueta.italic = True  # Aplica cursiva al prefijo.
        r_texto = p_nota.add_run(nota)  # Inserta la explicación.
        for run in (r_etiqueta, r_texto):  # Recorre ambas corridas de la nota.
            run.font.name = "Times New Roman"  # Mantiene la tipografía.
            run.font.size = Pt(9)  # Usa tamaño de nota legible.
    return t  # Devuelve la tabla construida.


def figura(numero, titulo, ruta, ancho, nota):  # Inserta una figura científica con rótulo APA.
    p_numero = documento.add_paragraph()  # Crea el número de figura.
    p_numero.paragraph_format.keep_with_next = True  # Mantiene el número junto al título.
    r_numero = p_numero.add_run(f"Figura {numero}")  # Inserta el rótulo numerado.
    r_numero.bold = True  # Resalta el número de figura.
    r_numero.font.name = "Times New Roman"  # Mantiene la tipografía.
    r_numero.font.size = Pt(11)  # Usa tamaño legible.
    p_titulo = documento.add_paragraph()  # Crea el título de figura.
    p_titulo.paragraph_format.keep_with_next = True  # Evita separar el título de la imagen.
    r_titulo = p_titulo.add_run(titulo)  # Inserta el título descriptivo.
    r_titulo.italic = True  # Sigue la convención APA.
    r_titulo.font.name = "Times New Roman"  # Mantiene la tipografía.
    r_titulo.font.size = Pt(11)  # Usa tamaño coherente con las tablas.
    p_imagen = documento.add_paragraph()  # Crea el contenedor de la imagen.
    p_imagen.alignment = WD_ALIGN_PARAGRAPH.CENTER  # Centra la figura en la página.
    p_imagen.paragraph_format.keep_with_next = True  # Mantiene la imagen junto a su nota.
    p_imagen.add_run().add_picture(str(ruta), width=Inches(ancho))  # Inserta la imagen validada.
    p_nota = documento.add_paragraph()  # Crea la nota de figura.
    p_nota.paragraph_format.line_spacing = 1  # Usa espacio sencillo para la nota.
    r_etiqueta = p_nota.add_run("Nota. ")  # Inserta el prefijo APA.
    r_etiqueta.italic = True  # Aplica cursiva al prefijo.
    r_texto = p_nota.add_run(nota)  # Inserta la fuente y la interpretación necesaria.
    for run in (r_etiqueta, r_texto):  # Recorre las corridas de la nota.
        run.font.name = "Times New Roman"  # Mantiene la tipografía.
        run.font.size = Pt(9)  # Usa un tamaño compacto.


documento = Document()  # Crea un documento Word vacío.
seccion = documento.sections[0]  # Selecciona la primera sección.
seccion.top_margin = Inches(1)  # Aplica margen superior de una pulgada.
seccion.bottom_margin = Inches(1)  # Aplica margen inferior de una pulgada.
seccion.left_margin = Inches(1)  # Aplica margen izquierdo de una pulgada.
seccion.right_margin = Inches(1)  # Aplica margen derecho de una pulgada.
agregar_numero_pagina(seccion.header.paragraphs[0])  # Añade numeración automática desde la portada.

estilo_normal = documento.styles["Normal"]  # Recupera el estilo base del documento.
estilo_normal.font.name = "Times New Roman"  # Define la tipografía general.
estilo_normal.font.size = Pt(12)  # Define el tamaño general.
estilo_normal.font.color.rgb = NEGRO  # Conserva el texto negro.
estilo_titulo = documento.styles["Title"]  # Recupera el estilo semántico usado en la portada.
propiedades_estilo_titulo = estilo_titulo.element.get_or_add_pPr()  # Obtiene las propiedades XML del estilo de título.
borde_estilo_titulo = propiedades_estilo_titulo.find(qn("w:pBdr"))  # Busca la línea decorativa heredada del tema de Word.
if borde_estilo_titulo is not None:  # Comprueba si el tema incluyó ese borde inferior.
    propiedades_estilo_titulo.remove(borde_estilo_titulo)  # Suprime la línea para cumplir la composición limpia de la portada.

documento.add_picture(str(ACTIVOS / "logo_uce_oficial.png"), width=Inches(1.35))  # Inserta el sello oficial en la portada.
documento.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER  # Centra el sello institucional.
documento.add_paragraph()  # Añade espacio visual antes del título.
p_titulo = documento.add_paragraph(style="Title")  # Crea el título principal con el estilo semántico de Word.
p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER  # Centra el título en la portada.
propiedades_titulo = p_titulo._p.get_or_add_pPr()  # Obtiene las propiedades XML del título de portada.
borde_titulo = propiedades_titulo.find(qn("w:pBdr"))  # Busca el borde decorativo heredado por el estilo de Word.
if borde_titulo is not None:  # Comprueba si el borde fue materializado directamente en el párrafo.
    propiedades_titulo.remove(borde_titulo)  # Elimina la línea para conservar una portada APA limpia.
r_titulo = p_titulo.add_run("Brechas de empleo adecuado en personas con educación superior completada en Ecuador")  # Inserta el título específico.
r_titulo.bold = True  # Destaca el título principal.
r_titulo.font.name = "Times New Roman"  # Usa tipografía APA.
r_titulo.font.size = Pt(16)  # Aplica el tamaño recomendado para portada.
r_titulo.font.color.rgb = NEGRO  # Mantiene el título en negro.
parrafo("Análisis estadístico y aprendizaje automático interpretable con ENEMDU 2025", cursiva=True, centrado=True, sangria=False)  # Añade el subtítulo metodológico.
documento.add_paragraph()  # Añade espacio antes de los datos académicos.
parrafo("Franz Eduardo Del Pozo Sánchez", centrado=True, sangria=False)  # Identifica al autor a partir de la fuente institucional verificada.
parrafo("Facultad de Ciencias, Universidad Central del Ecuador", centrado=True, sangria=False)  # Registra la afiliación institucional.
parrafo("Asignatura: Python para Ciencia de Datos e Inteligencia Artificial", centrado=True, sangria=False)  # Identifica la asignatura del trabajo final.
parrafo("Quito, Ecuador", centrado=True, sangria=False)  # Registra la localización institucional.
parrafo("1 de octubre de 2026", centrado=True, sangria=False)  # Registra la fecha de cierre de esta versión.
documento.add_page_break()  # Inicia el resumen en una página independiente.

encabezado("Resumen", 1)  # Abre el resumen en español.
parrafo("Este estudio cuantitativo transversal analizó las brechas de empleo adecuado entre personas económicamente activas de 15 años o más que declararon haber obtenido un título superior en Ecuador durante 2025. Se utilizaron los microdatos anuales de la Encuesta Nacional de Empleo, Desempleo y Subempleo, con factores de expansión, estratos y unidades primarias de muestreo. La población analítica descriptiva reunió 39 551 registros. La proporción ponderada de empleo adecuado fue 68,85% (IC 95% [67,54%, 70,13%]). Las mujeres presentaron una diferencia de −5,91 puntos porcentuales respecto de los hombres y el área rural una diferencia de −10,97 puntos respecto del área urbana. En la regresión logística ajustada, las razones de momios fueron 0,73 para mujeres, 0,80 para área rural, 0,86 para educación superior no universitaria y 2,88 para posgrado. La prueba reservada incluyó 7 791 registros agrupados en 1 269 UPM. El boosting alcanzó ROC AUC de 0,684 y la regresión logística 0,674; la diferencia pareada de 0,0106 tuvo un intervalo de confianza que incluyó cero. La logística mostró mejor pendiente de calibración y se retuvo como referencia. Persistieron brechas de sensibilidad por sexo y área. Los hallazgos describen asociaciones y clasificación contemporánea; no estiman efectos causales ni predicen empleo futuro.")  # Resume diseño, resultados y alcance sin exagerar conclusiones.
parrafo("Palabras clave: empleo adecuado, educación superior, encuesta compleja, aprendizaje automático, calibración, equidad algorítmica", cursiva=True, sangria=False)  # Añade palabras clave del estudio.

encabezado("Introducción", 1)  # Abre la introducción del informe.
parrafo("La obtención de un título superior no garantiza una inserción laboral adecuada. En Ecuador, la ENEMDU permite examinar esa relación con cobertura nacional, variables sociodemográficas, información territorial y un diseño muestral complejo. El problema exige dos lecturas complementarias. La inferencia describe diferencias poblacionales y su incertidumbre; la predicción evalúa cuánto permiten las características observadas clasificar la situación laboral de registros no utilizados durante el desarrollo del modelo.")  # Presenta el problema y la lógica dual del análisis.
parrafo("El proyecto responde a una pregunta aplicada: qué características se asocian con el empleo adecuado entre personas económicamente activas con educación superior completada y cuánto mejora la clasificación al usar un modelo de árboles frente a una regresión logística. La comparación considera discriminación, precisión, calibración y desempeño por grupos. El análisis evita usar ingresos, horas, condición de actividad u otras variables que intervienen directamente en la definición del resultado.")  # Explica la pregunta y la protección contra fuga de información.
parrafo("El informe sigue APA 7 para redacción y referencias. La parte observacional adopta STROBE como guía de transparencia para estudios transversales. La parte predictiva adapta TRIPOD+AI, aunque esa guía nació en investigación clínica, porque ofrece criterios útiles para describir datos, particiones, modelos, desempeño, calibración, subgrupos y ciencia abierta. PROBAST+AI se utiliza como marco crítico de riesgo de sesgo, no como certificación formal del modelo (Collins et al., 2024; Moons et al., 2025; von Elm et al., 2007).")  # Declara los estándares metodológicos recomendados.

encabezado("Problema de investigación", 1)  # Presenta el problema formal.
parrafo("La media nacional puede ocultar desigualdades dentro de la población titulada. Diferencias por sexo, área y territorio pueden persistir después de considerar edad, nivel educativo y otras características. Al mismo tiempo, un algoritmo flexible puede mejorar una métrica promedio y conservar errores desiguales entre grupos. El estudio separa estas preguntas para impedir que la capacidad predictiva se interprete como explicación causal o como evidencia automática de justicia.")  # Delimita la brecha de conocimiento.
encabezado("Pregunta principal", 2)  # Introduce la pregunta central.
parrafo("¿Qué características se asocian con tener empleo adecuado entre las personas económicamente activas con educación superior completada en Ecuador durante 2025, y cómo se compara el desempeño de un modelo de aprendizaje automático con el de una regresión logística para clasificar esa situación?", sangria=False)  # Expone la pregunta investigable.
encabezado("Objetivos", 2)  # Abre los objetivos del estudio.
parrafo("El objetivo general fue analizar brechas y asociaciones sociodemográficas y territoriales del empleo adecuado y comparar el desempeño fuera de muestra de una regresión logística con un modelo de boosting.")  # Declara el objetivo general.
tabla(("Tabla 1", "Objetivos específicos y evidencia producida"), ["Objetivo", "Evidencia"], [
    ("Estimar la proporción de empleo adecuado y sus diferencias por sexo, área, edad, educación y territorio.", "Tasas ponderadas, IC 95% y contrastes ajustados al diseño."),
    ("Estimar asociaciones ajustadas con variables preespecificadas.", "Razones de momios e IC 95% de una regresión logística con varianza de encuesta."),
    ("Comparar modelos sin contaminar la prueba.", "Cinco pliegues internos por UPM y evaluación única en la partición reservada."),
    ("Examinar calibración y errores por grupos.", "ROC AUC, PR AUC, Brier, sensibilidad, especificidad y brechas por sexo y área."),
    ("Comunicar resultados reproducibles.", "Código, cuadernos, bases tratadas, informes y dashboard sin identificadores."),
], [3.25, 3.25], "Los productos se vinculan con las etapas auditadas del repositorio reproducible.")  # Resume objetivos y productos.
encabezado("Hipótesis", 2)  # Abre las hipótesis preespecificadas.
parrafo("H1: la proporción de empleo adecuado difiere por sexo y territorio. H2: algunas diferencias persisten después del ajuste por edad, educación y covariables pertinentes. H3: el modelo de árboles mejora la clasificación fuera de muestra frente a la regresión logística. La tercera hipótesis podía rechazarse; el diseño no asumió de antemano una ventaja del aprendizaje automático.")  # Declara las hipótesis y su carácter refutable.

encabezado("Método", 1)  # Abre la sección metodológica principal.
encabezado("Diseño y fuente de datos", 2)  # Describe el diseño.
parrafo("Se realizó un estudio observacional transversal con análisis secundario de la base anual ENEMDU 2025. El archivo original de personas contiene 334 786 registros y 139 columnas correspondientes a los doce meses. La base anual representa observaciones de levantamientos mensuales y no debe interpretarse como una cohorte de personas únicas. La documentación oficial utilizada incluyó la guía de uso, el diseño muestral, el diccionario, los formularios mensuales, el manual del encuestador y la sintaxis de indicadores del INEC (Instituto Nacional de Estadística y Censos [INEC], 2026a, 2026b).")  # Identifica diseño, alcance y documentación.
encabezado("Población y criterios de elegibilidad", 2)  # Define la población analítica.
parrafo("La población principal comprendió registros de personas de 15 años o más, pertenecientes a la población económicamente activa, con nivel de instrucción superior no universitario, superior universitario o posgrado y respuesta afirmativa a la obtención de algún título superior. La regla exigió p03 ≥ 15; p10a en 8, 9 o 10; p12a = 1; y condact entre 1 y 8. Esta definición identifica titulación declarada y no verificación administrativa de credenciales.")  # Especifica la regla reproducible.
tabla(("Tabla 2", "Flujo analítico y partición predictiva"), ["Etapa", "Registros", "UPM", "Interpretación"], [
    ("Base original de personas", "334 786", "—", "Registros de los doce levantamientos de 2025."),
    ("Dominio descriptivo", "39 551", "6 326", "PEA con título superior declarado."),
    ("Dominio predictivo", "39 293", "6 323", "Excluye empleo no clasificado del resultado principal."),
    ("Desarrollo", "31 502", "5 054", "Cinco pliegues internos agrupados por UPM."),
    ("Prueba reservada", "7 791", "1 269", "Evaluada una sola vez tras congelar modelos y umbrales."),
], [2.0, 1.0, 0.9, 2.6], "Las claves de vivienda y hogar no se repartieron entre desarrollo y prueba; la UPM fue la unidad de agrupación principal.")  # Documenta el flujo de participantes.
encabezado("Variables", 2)  # Abre la operacionalización.
parrafo("El resultado descriptivo tomó valor uno cuando condact = 1, empleo adecuado o pleno, y cero para las demás categorías de la PEA elegible. El resultado predictivo excluyó condact = 6, empleo no clasificado, porque esa categoría no demuestra una situación laboral inadecuada. Los predictores autorizados fueron edad, indicador de 98 años o más, sexo, estado civil, asistencia educativa, nivel educativo reportado, autoidentificación étnica, provincia, área y mes.")  # Define resultado y predictores.
parrafo("Ingresos, horas, condición de actividad, variables derivadas del objetivo, identificadores y variables del diseño se excluyeron de los predictores. Esta lista permitida evitó fuga de información. Los factores de expansión se utilizaron en estimación y evaluación ponderada, mientras estrato y UPM definieron la incertidumbre y las particiones; ninguna de estas variables actuó como predictor.")  # Explica exclusiones y papeles de diseño.
encabezado("Calidad y tratamiento", 2)  # Resume controles de calidad.
parrafo("La auditoría verificó integridad del ZIP, correspondencia de 139 columnas, cobertura de los doce meses, códigos educativos en los formularios, valores especiales de edad e ingresos, factores positivos, relación entre estratos y UPM, identificadores y repeticiones entre levantamientos. No se eliminaron observaciones por compartir una clave de persona sin mes, pues la estructura anual no acredita identidad longitudinal. La base tratada conservó trazabilidad, etiquetas y bitácora de transformaciones.")  # Resume la auditoría sin confundir repetición con duplicación.
encabezado("Análisis exploratorio", 2)  # Describe las siete etapas EDA.
parrafo("El análisis exploratorio siguió siete componentes: comprensión y delimitación; calidad e identificadores; tratamiento y partición; análisis univariado; análisis bivariado; inferencia con diseño complejo; y desarrollo de modelos interpretables. Las distribuciones y frecuencias se calcularon con pesos cuando correspondía. Los gráficos mostraron denominadores y evitaron inferencias a partir de celdas pequeñas.")  # Documenta el flujo de EDA.
encabezado("Inferencia con diseño muestral", 2)  # Explica estimación e incertidumbre.
parrafo("Las proporciones se estimaron como cocientes ponderados. La linealización de Taylor incorporó estratos y UPM para obtener errores estándar e intervalos de confianza del 95%. Las diferencias entre grupos se expresaron en puntos porcentuales. La regresión logística asociativa usó una especificación predefinida y errores estándar del diseño. Las razones de momios describen asociaciones condicionales y no efectos causales.")  # Explica inferencia y su interpretación.
encabezado("Desarrollo y evaluación predictiva", 2)  # Describe el protocolo de modelado.
parrafo("La separación desarrollo-prueba se realizó por UPM y se comprobó la ausencia de solapamiento de viviendas y hogares. En desarrollo se usaron cinco pliegues internos también agrupados. Se compararon un clasificador base, una regresión logística y gradient boosting. Las transformaciones se ajustaron dentro de cada pliegue. Los hiperparámetros y umbrales se congelaron antes de abrir la prueba.")  # Documenta prevención de contaminación.
parrafo("La evaluación única en prueba informó ROC AUC, PR AUC, Brier, pérdida logarítmica, sensibilidad, especificidad, precisión, exactitud y exactitud balanceada. Los intervalos se obtuvieron con 1 000 réplicas bootstrap de UPM dentro de estrato. La comparación de modelos fue pareada. La calibración se examinó mediante intercepto, pendiente y deciles. El desempeño por sexo y área se reportó con el mismo umbral global, sin crear reglas diferentes por grupo.")  # Explica métricas y subgrupos.
encabezado("Ética, privacidad y reproducibilidad", 2)  # Presenta salvaguardas.
parrafo("El estudio utilizó microdatos públicos anonimizados. El dashboard publica resultados agregados y omite identificadores de persona, hogar y vivienda. No produce predicciones individuales ni recomendaciones laborales. Los archivos originales se conservaron sin cambios, las transformaciones quedaron registradas y la evaluación de prueba se cerró mediante huellas de archivos para impedir una segunda selección oportunista.")  # Explica ética y ciencia abierta.

encabezado("Resultados", 1)  # Abre la sección de resultados.
encabezado("Descripción e inferencia", 2)  # Presenta resultados poblacionales.
parrafo("La proporción ponderada de empleo adecuado en el dominio fue 68,85% (IC 95% [67,54%, 70,13%]). Los hombres alcanzaron 72,10% y las mujeres 66,19%. El área urbana alcanzó 70,41% y el área rural 59,44%. Por educación, el posgrado registró 85,58%, la educación superior universitaria 66,28% y la superior no universitaria 63,91%.")  # Resume estimaciones principales.
figura(1, "Proporciones de empleo adecuado por sexo y área", FIGURAS / "eda06_inferencia" / "06_ic_contrastes_primarios.png", 6.2, "Estimaciones ponderadas e intervalos de confianza del 95% que incorporan estrato y UPM. Elaboración propia con ENEMDU anual 2025.")  # Inserta los contrastes principales.
tabla(("Tabla 3", "Empleo adecuado en el dominio de educación superior completada"), ["Grupo", "n", "%", "IC 95%", "CV (%)"], [
    ("Total", "39 551", "68,85", "[67,54; 70,13]", "0,96"),
    ("Hombre", "18 198", "72,10", "[70,41; 73,73]", "1,18"),
    ("Mujer", "21 353", "66,19", "[64,51; 67,83]", "1,28"),
    ("Urbana", "35 099", "70,41", "[69,04; 71,74]", "0,98"),
    ("Rural", "4 452", "59,44", "[55,70; 63,07]", "3,17"),
    ("Superior no universitaria", "6 425", "63,91", "[61,17; 66,56]", "2,15"),
    ("Superior universitaria", "25 864", "66,28", "[64,68; 67,84]", "1,21"),
    ("Posgrado", "7 262", "85,58", "[83,43; 87,50]", "1,21"),
], [2.4, 0.8, 0.7, 1.5, 0.8], "n es el tamaño sin ponderar. CV es el coeficiente de variación de la proporción estimada.", mantener_junta=True)  # Presenta la tabla científica principal sin dividir sus filas entre páginas.
parrafo("La diferencia mujeres menos hombres fue −5,91 puntos porcentuales y la diferencia rural menos urbana fue −10,97 puntos. Ambas respaldan la primera hipótesis descriptiva. La magnitud territorial fue mayor, aunque la precisión rural resultó menor por su tamaño efectivo y estructura muestral.")  # Interpreta los contrastes.
encabezado("Asociaciones ajustadas", 2)  # Presenta la regresión asociativa.
figura(2, "Asociaciones ajustadas seleccionadas con empleo adecuado", FIGURAS / "eda07_modelado" / "04_asociaciones_ajustadas.png", 6.1, "Razones de momios e intervalos del 95% de la regresión logística asociativa. Las categorías de referencia y la especificación completa constan en los archivos del proyecto. Elaboración propia.")  # Inserta el bosque de asociaciones.
tabla(("Tabla 4", "Asociaciones ajustadas principales"), ["Contraste", "OR", "IC 95%", "p"], [
    ("Mujer frente a hombre", "0,73", "[0,66; 0,81]", "< ,001"),
    ("Rural frente a urbana", "0,80", "[0,68; 0,94]", ",007"),
    ("Superior no universitaria frente a universitaria", "0,86", "[0,76; 0,98]", ",029"),
    ("Posgrado frente a universitaria", "2,88", "[2,42; 3,43]", "< ,001"),
], [3.2, 0.8, 1.3, 0.7], "OR es la razón de momios. El modelo ajustó además por edad no lineal, estado civil, asistencia educativa, autoidentificación étnica, provincia y mes.")  # Resume asociaciones clave.
parrafo("Las asociaciones ajustadas conservaron desventajas para mujeres y residentes rurales. El posgrado mostró una asociación positiva marcada. Estos resultados apoyan la segunda hipótesis, pero la simultaneidad de medición y la posibilidad de confusión residual impiden interpretarlos como efectos de sexo, territorio o educación.")  # Interpreta sin causalidad.
encabezado("Desempeño fuera de muestra", 2)  # Presenta la prueba reservada.
figura(3, "Discriminación en la prueba reservada", FIGURAS / "eda08_evaluacion" / "01_curvas_roc_pr.png", 6.2, "Curvas ROC y precisión-recall ponderadas para la regresión logística y el boosting en 7 791 registros reservados. Elaboración propia.")  # Inserta la comparación predictiva.
tabla(("Tabla 5", "Métricas principales en la prueba reservada"), ["Modelo", "ROC AUC", "PR AUC", "Brier", "Exactitud balanceada"], [
    ("Base", "0,500", "0,687", "0,215", "0,500"),
    ("Logística", "0,674 [0,645; 0,702]", "0,813 [0,783; 0,842]", "0,198 [0,185; 0,210]", "0,632 [0,604; 0,659]"),
    ("Boosting", "0,684 [0,654; 0,714]", "0,824 [0,794; 0,852]", "0,195 [0,181; 0,210]", "0,621 [0,592; 0,647]"),
], [1.1, 1.45, 1.45, 1.45, 1.4], "Los intervalos del 95% provienen de 1 000 réplicas bootstrap de UPM dentro de estrato. Menor Brier indica mejor precisión probabilística.")  # Resume las métricas globales.
parrafo("El boosting mejoró el ROC AUC en 0,0106 respecto de la logística, pero el IC 95% de la diferencia [−0,0097; 0,0303] incluyó cero. La diferencia de PR AUC fue 0,0119 [−0,0012; 0,0237]. Por tanto, la tercera hipótesis no recibió apoyo concluyente en la prueba reservada.")  # Responde la hipótesis predictiva.
figura(4, "Calibración de los modelos en la prueba reservada", FIGURAS / "eda08_evaluacion" / "02_calibracion_prueba.png", 6.0, "Probabilidad observada frente a predicha por deciles. La línea diagonal representa calibración ideal. La pendiente fue 0,930 para logística y 0,761 para boosting. Elaboración propia.")  # Inserta el análisis de calibración.
parrafo("La regresión logística presentó una pendiente de calibración de 0,930, más cercana a uno que la pendiente de 0,761 del boosting. Esta diferencia, sumada a la ausencia de una mejora concluyente en discriminación, justificó seleccionar la logística como modelo principal de referencia y conservar el boosting como comparador.")  # Explica la decisión de modelo.
encabezado("Desempeño por grupos", 2)  # Presenta la evaluación de equidad.
figura(5, "Brechas de sensibilidad por sexo y área", FIGURAS / "eda08_evaluacion" / "04_brechas_sensibilidad_ic95.png", 6.0, "Diferencias de sensibilidad con el umbral global congelado. Valores negativos indican menor sensibilidad para mujeres o para el área rural. Elaboración propia.")  # Inserta el gráfico de brechas.
parrafo("En boosting, la sensibilidad fue 69,2% en hombres y 53,9% en mujeres, con una brecha de −15,3 puntos [−20,5; −9,7]. La sensibilidad fue 62,5% en el área urbana y 53,5% en el área rural, con una brecha de −9,0 puntos [−17,7; −0,7]. La ROC AUC rural fue mayor que la urbana, lo que demuestra que una buena discriminación no garantiza un punto de operación equivalente entre grupos.")  # Interpreta la coexistencia de métricas.

encabezado("Discusión", 1)  # Abre la discusión científica.
parrafo("Los resultados muestran brechas de empleo adecuado incluso dentro de una población con título superior declarado. La diferencia por sexo persiste después del ajuste y la desventaja rural aparece tanto en proporciones como en el modelo asociativo. El posgrado se relaciona con mayores momios de empleo adecuado, aunque el estudio no puede separar selección educativa, campo de formación, experiencia, estructura productiva u otras causas.")  # Resume el hallazgo sustantivo.
parrafo("La comparación predictiva aporta una conclusión práctica. El boosting produjo mejores valores puntuales de ROC AUC, PR AUC y Brier, pero las diferencias pareadas fueron pequeñas e inciertas. La logística resultó competitiva y mejor calibrada. Para este conjunto de predictores sociodemográficos, la flexibilidad adicional del modelo de árboles no justificó presentarlo como superior.")  # Interpreta el contraste entre modelos.
parrafo("La evaluación por grupos modifica la lectura de las métricas globales. Ambos modelos tuvieron menor sensibilidad para mujeres y residentes rurales con un umbral común. Este resultado no demuestra discriminación causal del algoritmo, pero sí indica que una eventual aplicación debe estudiar objetivos, costos de error y condiciones de uso antes de fijar umbrales. Ajustar umbrales por grupo sería una decisión normativa que excede este análisis.")  # Discute equidad sin sobreafirmar.
encabezado("Limitaciones", 2)  # Expone límites materiales.
parrafo("El diseño transversal impide establecer temporalidad y causalidad. La titulación es autodeclarada y no identifica de manera armonizada el campo de estudio ni la calidad institucional. La base anual agrega levantamientos mensuales y no constituye una cohorte. La condición de empleo adecuado responde a una clasificación contemporánea, por lo que los modelos no pronostican empleo futuro. La prueba reservada ofrece validación interna y no validación temporal o externa. Los intervalos de desempeño son condicionales a una partición y el bootstrap aproxima la variación por conglomerados.")  # Enumera limitaciones que afectan interpretación.
encabezado("Conclusiones", 1)  # Abre las conclusiones.
parrafo("La proporción nacional estimada de empleo adecuado entre personas económicamente activas con título superior declarado fue cercana a siete de cada diez. Las brechas por sexo y área fueron estadísticamente precisas y persistieron tras el ajuste. El boosting no mostró una ventaja concluyente frente a la regresión logística en prueba, y su calibración fue inferior. La regresión logística constituye la referencia más prudente para comunicación e interpretación, mientras el boosting permite examinar posibles no linealidades y diferencias de error.")  # Sintetiza la respuesta principal.
parrafo("El dashboard complementa el informe al permitir explorar tasas, intervalos, tamaños muestrales, asociaciones y métricas congeladas. Su uso debe permanecer descriptivo y científico. No debe utilizarse para perfilar personas, asignar oportunidades ni sustituir evaluaciones laborales.")  # Delimita el uso del producto aplicado.
encabezado("Recomendaciones para el artículo", 1)  # Presenta la ruta de publicación.
parrafo("El manuscrito publicable puede concentrarse en dos contribuciones: la persistencia de brechas de empleo adecuado dentro de la población titulada y la ausencia de una ventaja predictiva concluyente del boosting cuando se exige calibración y evaluación por grupos. Antes del envío conviene añadir validación temporal con otro año, análisis de sensibilidad de la definición educativa, revisión completa con las listas STROBE y TRIPOD+AI, evaluación PROBAST+AI adaptada al contexto no clínico y un plan público de código y materiales sin microdatos restringidos.")  # Define pasos concretos para publicación.

documento.add_page_break()  # Inicia las referencias en una página independiente.
encabezado("Referencias", 1)  # Abre la lista de referencias APA.
referencias = [  # Reúne las fuentes citadas en orden alfabético.
    "American Psychological Association. (2020). Publication manual of the American Psychological Association (7th ed.).",
    "Collins, G. S., Moons, K. G. M., Dhiman, P., Riley, R. D., Beam, A. L., Van Calster, B., Ghassemi, M., Liu, X., Reitsma, J. B., van Smeden, M., Boulesteix, A. L., Camaradou, J. C., Celi, L. A., Denaxas, S., Denniston, A. K., Glocker, B., Golub, R. M., Harvey, H., Heinze, G., ... Logullo, P. (2024). TRIPOD+AI statement: Updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ, 385, e078378. https://doi.org/10.1136/bmj-2023-078378",
    "Instituto Nacional de Estadística y Censos. (2026a). Diseño muestral de la Encuesta Nacional de Empleo, Desempleo y Subempleo anual enero-diciembre 2025. https://www.ecuadorencifras.gob.ec/documentos/web-inec/EMPLEO/2025/anual/Disenio_Muestral_ENEMDU_Anual_enero-diciembre_2025.pdf",
    "Instituto Nacional de Estadística y Censos. (2026b). Guía de uso de base de datos ENEMDU 2021-2025. https://www.ecuadorencifras.gob.ec/documentos/web-inec/EMPLEO/2025/anual/Guia_de_usuario_BDD_ENEMDU_anual_2025.pdf",
    "Instituto Nacional de Estadística y Censos. (2026c). Microdatos ENEMDU anual 2025 [Conjunto de datos]. https://www.ecuadorencifras.gob.ec/documentos/web-inec/EMPLEO/2025/anual/2_BDD_DATOS_ABIERTOS_ENEMDU_2025_CSV.zip",
    "Moons, K. G. M., Damen, J. A. A., Kaul, T., Hooft, L., Andaur Navarro, C., Dhiman, P., Beam, A. L., Van Calster, B., Celi, L. A., Denaxas, S., ... Collins, G. S. (2025). PROBAST+AI: An updated quality, risk of bias, and applicability assessment tool for prediction models using regression or artificial intelligence methods. BMJ, 388, e082505. https://doi.org/10.1136/bmj-2024-082505",
    "von Elm, E., Altman, D. G., Egger, M., Pocock, S. J., Gøtzsche, P. C., & Vandenbroucke, J. P. (2007). The Strengthening the Reporting of Observational Studies in Epidemiology (STROBE) statement: Guidelines for reporting observational studies. PLoS Medicine, 4(10), e296. https://doi.org/10.1371/journal.pmed.0040296",
]  # Cierra la colección de referencias.
for referencia in referencias:  # Recorre las referencias en orden alfabético.
    p = documento.add_paragraph()  # Crea un párrafo por referencia.
    p.paragraph_format.left_indent = Inches(0.5)  # Desplaza el bloque de referencia.
    p.paragraph_format.first_line_indent = Inches(-0.5)  # Aplica sangría francesa APA.
    p.paragraph_format.line_spacing = 2  # Mantiene doble espacio.
    p.paragraph_format.space_after = Pt(0)  # Evita separación adicional.
    r = p.add_run(referencia)  # Inserta la referencia completa.
    r.font.name = "Times New Roman"  # Usa tipografía APA.
    r.font.size = Pt(12)  # Usa tamaño estándar.

documento.add_page_break()  # Inicia los anexos en página independiente.
encabezado("Apéndice A Matriz resumida de variables", 1)  # Abre el apéndice de operacionalización.
tabla(("Tabla A1", "Variables principales y función analítica"), ["Variable", "Código", "Función", "Regla"], [
    ("Edad", "p03", "Selección y predictor", "15 o más; 98 indica 98 años o más."),
    ("Sexo registrado", "p02", "Comparación principal", "Categorías oficiales hombre y mujer."),
    ("Nivel educativo", "p10a", "Selección y predictor", "8 superior no universitaria, 9 universitaria, 10 posgrado."),
    ("Título superior", "p12a", "Selección", "Exigir respuesta sí."),
    ("Condición de actividad", "condact", "Selección y resultado", "1 adecuado; 2–5 y 7–8 no adecuado en predicción."),
    ("Factor de expansión", "fexp", "Ponderación", "Solo estimación y métricas, nunca predictor."),
    ("Estrato y UPM", "estrato, upm", "Varianza y partición", "Conservar estructura de diseño y agrupación."),
], [1.35, 1.05, 1.55, 2.55], "La matriz completa y su trazabilidad se encuentran en docs/03_matriz_inicial_variables.md.")  # Resume la operacionalización.
encabezado("Apéndice B Productos reproducibles", 1)  # Abre el inventario de archivos.
parrafo("El repositorio contiene los programas Python por etapa, cuadernos ejecutables, reportes de calidad, base tratada, diccionario, modelos congelados, predicciones de prueba cerradas, tablas científicas, figuras SVG y PNG, y el dashboard Streamlit. El archivo README.md documenta el orden de ejecución. Los datos originales permanecen en data/raw y los productos del dashboard en data/dashboard.")  # Documenta la estructura de reproducibilidad.
parrafo("Declaración de uso de herramientas: la programación y la edición del presente informe se apoyaron en herramientas de asistencia automatizada. Las decisiones analíticas, las definiciones operativas, la revisión de fuentes, la evaluación de resultados y la responsabilidad final corresponden al autor.")  # Añade una declaración transparente de asistencia.

documento.save(SALIDA)  # Guarda el documento Word editable en la carpeta final.
print(SALIDA)  # Informa la ruta creada para la etapa de verificación.
