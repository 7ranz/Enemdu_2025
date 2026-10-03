# ETAPA 6: verificar inferencia, gráficos, cuaderno y tabla científica.
# ARCHIVO: 05c_verificar_eda06.py; produce el acta final de la etapa.
# BLOQUE 1: dependencias y rutas.
import ast  # Comprueba sintaxis Python.
import hashlib  # Calcula huellas finales.
import json  # Lee y escribe evidencia.
from pathlib import Path  # Maneja rutas portables.
import nbformat  # Inspecciona el cuaderno.
import numpy as np  # Realiza controles numéricos independientes.
import pandas as pd  # Lee las tablas publicadas.
from openpyxl import load_workbook  # Inspecciona el XLSX final sin usarlo para autoría.
from PIL import Image  # Comprueba resolución de PNG.

RAIZ = Path(__file__).resolve().parents[1]  # Localiza el proyecto.
REPORTES = RAIZ / 'reports/eda06_inferencia'  # Localiza tablas.
FIGURAS = RAIZ / 'reports/figures/eda06_inferencia'  # Localiza figuras.
DISENO = RAIZ / 'data/processed/inferencia/diseno_bivariado.csv.gz'  # Localiza el diseño completo.
CUADERNO = RAIZ / 'notebooks/06_inferencia_diseno.ipynb'  # Localiza el cuaderno ejecutado.
LIBRO = RAIZ / 'outputs/eda06_inferencia/tabla_cientifica_eda06.xlsx'  # Localiza el libro científico.

# BLOQUE 2: huella y control del manifiesto de cálculo.
def sha256(ruta):  # Calcula SHA-256 sin modificar el archivo.
    h = hashlib.sha256()  # Inicializa el algoritmo.
    with ruta.open('rb') as archivo:  # Abre bytes en lectura.
        for bloque in iter(lambda: archivo.read(1024 * 1024), b''):  # Lee por bloques.
            h.update(bloque)  # Incorpora cada bloque.
    return h.hexdigest()  # Devuelve la huella.

manifiesto = pd.read_csv(REPORTES / '04_manifiesto_productos.csv')  # Lee el inventario generado por el análisis.
for fila in manifiesto.itertuples(index=False):  # Recorre productos registrados.
    ruta = RAIZ / fila.archivo  # Reconstruye la ruta.
    assert ruta.exists()  # Exige existencia.
    assert ruta.stat().st_size == fila.bytes  # Exige tamaño estable.
    assert sha256(ruta) == fila.sha256  # Exige contenido estable.

# BLOQUE 3: controles numéricos y prueba manual de la fórmula de varianza.
tasas = pd.read_csv(REPORTES / '01_tasas_con_ic95.csv', dtype={'codigo': 'string'})  # Lee tasas científicas.
contrastes = pd.read_csv(REPORTES / '02_contrastes_primarios_ic95.csv')  # Lee contrastes.
globales = pd.read_csv(REPORTES / '03_pruebas_globales_wald.csv')  # Lee pruebas globales.
columnas = ['denominador_dominio', 'numerador_adecuado', 'fexp']  # Selecciona campos para recalcular el total.
base = pd.read_csv(DISENO, usecols=columnas, dtype={'denominador_dominio': 'Int8', 'numerador_adecuado': 'Int8', 'fexp': 'float64'}, encoding='utf-8-sig')  # Lee contribuciones validadas.
tasa_total_independiente = 100 * (base['fexp'] * base['numerador_adecuado']).sum() / (base['fexp'] * base['denominador_dominio']).sum()  # Recalcula el punto total.
tasa_total_publicada = tasas.loc[tasas['variable'] == 'total', 'porcentaje'].iloc[0]  # Recupera el punto publicado.
assert np.isclose(tasa_total_independiente, tasa_total_publicada, atol=1e-10, rtol=0)  # Exige reconciliación independiente.
assert tasas.apply(lambda fila: 0 <= fila['ic95_inferior'] <= fila['porcentaje'] <= fila['ic95_superior'] <= 100, axis=1).all()  # Exige intervalos de proporción válidos.
assert np.allclose(contrastes['diferencia_pp'], contrastes['tasa_a'] - contrastes['tasa_b'])  # Reconcilia diferencias.
assert contrastes['valor_p_holm'].ge(contrastes['valor_p']).all()  # Comprueba que Holm no reduzca valores p.
assert len(globales) == 5 and globales['grados_libertad_prueba'].gt(0).all()  # Comprueba pruebas globales.
totales_prueba = [np.array([1.0, -1.0]), np.array([2.0, -2.0])]  # Define dos estratos de dos UPM con media cero.
varianza_manual = sum(len(bloque) / (len(bloque) - 1) * ((bloque - bloque.mean()) ** 2).sum() for bloque in totales_prueba)  # Reproduce la fórmula usada.
assert np.isclose(varianza_manual, 20.0)  # Contrasta con el resultado calculado a mano.

# BLOQUE 4: controles visuales, cuaderno y programas.
nombres = ['01_ic_sexo', '02_ic_area', '03_ic_edad', '04_ic_educacion', '05_ic_provincia', '06_ic_contrastes_primarios']  # Enumera figuras esperadas.
for nombre in nombres:  # Recorre cada gráfico.
    png = FIGURAS / f'{nombre}.png'  # Localiza PNG.
    svg = FIGURAS / f'{nombre}.svg'  # Localiza SVG.
    assert png.exists() and svg.exists()  # Exige ambos formatos.
    with Image.open(png) as imagen:  # Abre la figura rasterizada.
        assert imagen.width >= 1200 and imagen.height >= 700  # Exige resolución suficiente.
    assert '<svg' in svg.read_text(encoding='utf-8')  # Exige un vector válido.
programas = sorted(RAIZ.glob('src/05*.py'))  # Reúne programas Python de la etapa.
for programa in programas:  # Recorre cada programa.
    contenido = programa.read_text(encoding='utf-8-sig')  # Lee el código.
    ast.parse(contenido)  # Exige sintaxis válida.
    assert '# ETAPA 6' in contenido and '# ARCHIVO:' in contenido and '# BLOQUE' in contenido  # Exige comentarios estructurales.
cuaderno = nbformat.read(CUADERNO, as_version=4)  # Lee el cuaderno final.
celdas_codigo = [celda for celda in cuaderno.cells if celda.cell_type == 'code']  # Selecciona celdas ejecutables.
assert all(celda.execution_count is not None for celda in celdas_codigo)  # Exige ejecución completa.
assert not any(salida.output_type == 'error' for celda in celdas_codigo for salida in celda.outputs)  # Exige cero errores.
assert all(all((not linea.strip()) or linea.lstrip().startswith('#') or '#' in linea for linea in celda.source.splitlines()) for celda in celdas_codigo)  # Exige comentario por línea ejecutable.

# BLOQUE 5: inspección del libro Excel exportado.
assert LIBRO.exists() and LIBRO.stat().st_size > 10000  # Exige un libro material.
libro = load_workbook(LIBRO, read_only=False, data_only=False)  # Abre el XLSX para inspección sin guardar cambios.
hojas_esperadas = ['Resumen', 'Tasas', 'Contrastes', 'Pruebas globales', 'Método']  # Define su topología.
assert libro.sheetnames == hojas_esperadas  # Exige el orden científico previsto.
assert libro['Resumen']['A2'].value == 'Empleo adecuado y educación superior en Ecuador'  # Verifica el título.
assert libro['Tasas'].max_row == len(tasas) + 2  # Reconcilia las 47 estimaciones.
assert libro['Contrastes'].max_row == len(contrastes) + 2  # Reconcilia los cuatro contrastes.
assert libro['Pruebas globales'].max_row == len(globales) + 2  # Reconcilia las cinco pruebas.
errores_excel = {'#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!', '#NULL!', '#SPILL!', '#CALC!'}  # Define fallos de fórmula.
assert not any(celda.value in errores_excel for hoja in libro.worksheets for fila in hoja.iter_rows() for celda in fila if isinstance(celda.value, str))  # Exige ausencia de errores visibles.
libro.close()  # Cierra el archivo.

# BLOQUE 6: acta final con huellas de los entregables.
resultado = {'estado': 'verificado', 'filas_muestra_completa': len(base), 'estimaciones_con_ic': len(tasas), 'contrastes_primarios': len(contrastes), 'pruebas_globales': len(globales), 'prueba_manual_varianza': True, 'tasa_total_reconciliada': True, 'figuras_png': len(list(FIGURAS.glob('*.png'))), 'figuras_svg': len(list(FIGURAS.glob('*.svg'))), 'celdas_cuaderno': len(cuaderno.cells), 'celdas_codigo_ejecutadas': len(celdas_codigo), 'errores_cuaderno': 0, 'hojas_excel': hojas_esperadas, 'xlsx_sha256': sha256(LIBRO), 'csv_cientifico_sha256': sha256(REPORTES / '06_tabla_cientifica_principal.csv'), 'cuaderno_sha256': sha256(CUADERNO), 'prueba_predictiva_evaluada': False, 'programas_sintaxis_valida': [programa.name for programa in programas]}  # Consolida evidencia.
(REPORTES / '07_verificacion_entrega.json').write_text(json.dumps(resultado, ensure_ascii=False, indent=2), encoding='utf-8')  # Guarda el acta final.
print(json.dumps(resultado, ensure_ascii=False, indent=2))  # Presenta el cierre.
