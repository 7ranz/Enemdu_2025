# ETAPA 0: comprobar las fuentes antes de definir la base analítica.
# ARCHIVO: 00_verificar_fuentes.py. Lee originales y genera un reporte; no modifica datos.
# REQUISITOS: Python 3.10 o posterior y openpyxl para leer el diccionario Excel.

# BLOQUE 1: importar herramientas y localizar archivos con rutas portables.
import csv  # Lee tablas de texto respetando separadores y comillas.
import hashlib  # Calcula una huella que permite reconocer el archivo descargado.
import io  # Convierte los bytes del ZIP en texto legible.
import json  # Guarda resultados estructurados que podremos reutilizar.
from collections import Counter  # Cuenta categorías sin cargar toda la tabla en memoria.
from datetime import datetime, timezone  # Registra cuándo se ejecutó la comprobación.
from pathlib import Path  # Construye rutas válidas en distintos sistemas operativos.
from zipfile import ZipFile  # Permite leer archivos comprimidos sin alterar el original.
from openpyxl import load_workbook  # Abre el diccionario oficial en modo de solo lectura.

RAIZ = Path(__file__).resolve().parents[1]  # Localiza la carpeta enemdu_2025 desde este programa.
ORIGINALES = RAIZ / 'data' / 'raw'  # Señala la carpeta de fuentes originales.
SALIDAS = RAIZ / 'reports'  # Señala la carpeta donde se guardará la auditoría.
SALIDAS.mkdir(parents=True, exist_ok=True)  # Crea la carpeta de reportes si aún no existe.
ARCHIVO_ZIP = ORIGINALES / '2_BDD_DATOS_ABIERTOS_ENEMDU_2025_CSV.zip'  # Localiza los microdatos.
DICCIONARIO = ORIGINALES / 'documentacion' / 'diccionario' / 'Diccionario de Datos_persona_anual_2025.xlsx'  # Localiza el diccionario.

# BLOQUE 2: leer los nombres de variables publicados por el INEC.
libro = load_workbook(DICCIONARIO, read_only=True, data_only=True)  # Abre Excel sin editarlo.
hoja = libro['Hoja1']  # Selecciona la hoja que contiene nombres y descripciones.
variables_documentadas = {fila[0]: fila[1] for fila in hoja.iter_rows(min_row=7, values_only=True) if fila[0]}  # Conserva las filas con variable.
libro.close()  # Libera el archivo después de leerlo.

# BLOQUE 3: verificar integridad y recorrer los registros sin calcular asociaciones.
conteos_mes = Counter()  # Prepara el conteo de registros de cada mes.
conteos_educacion = Counter()  # Prepara el cruce estructural entre nivel y título.
conteos_condact = Counter()  # Prepara la revisión de códigos de condición de actividad.
candidatos_mes = Counter()  # Cuenta registros que cumplen la propuesta de elegibilidad.
faltantes_titulo_mes = Counter()  # Detecta no respuesta en el título entre niveles superiores.
edades_invalidas = 0  # Inicializa el número de edades vacías o no interpretables.
requeridas = {'p03', 'p10a', 'p12a', 'condact', 'mes', 'fexp', 'estrato', 'upm', 'id_hogar', 'id_persona'}  # Define campos indispensables.
with ZipFile(ARCHIVO_ZIP) as comprimido:  # Abre el contenedor oficial solo para lectura.
    miembro_danado = comprimido.testzip()  # Comprueba la integridad CRC de cada miembro.
    if miembro_danado is not None:  # Detiene el trabajo si el ZIP está dañado.
        raise ValueError(f'Archivo dañado dentro del ZIP: {miembro_danado}')  # Explica el problema.
    inventario = [{'archivo': item.filename, 'bytes': item.file_size} for item in comprimido.infolist()]  # Registra el contenido.
    with comprimido.open('BDDenemdu_personas_2025_anual.csv') as binario:  # Abre la base de personas.
        texto = io.TextIOWrapper(binario, encoding='utf-8-sig', newline='')  # Interpreta UTF-8 y retira la marca BOM.
        lector = csv.DictReader(texto, delimiter=';')  # Asocia cada valor con su encabezado.
        columnas = lector.fieldnames or []  # Conserva los nombres de columnas encontrados.
        ausentes = requeridas - set(columnas)  # Identifica campos necesarios que no estén presentes.
        if ausentes:  # Impide calcular con una estructura incompleta.
            raise ValueError(f'Faltan columnas necesarias: {sorted(ausentes)}')  # Enumera los campos ausentes.
        for fila in lector:  # Recorre los registros uno por uno sin modificarlos.
            mes = fila['mes'].strip()  # Lee el mes como categoría, no como resultado laboral.
            nivel = fila['p10a'].strip()  # Lee el código del nivel de instrucción.
            titulo = fila['p12a'].strip()  # Lee si se declaró algún título superior.
            condicion = fila['condact'].strip()  # Lee la categoría laboral oficial.
            conteos_mes[mes] += 1  # Suma un registro al mes correspondiente.
            conteos_educacion[(nivel, titulo)] += 1  # Cuenta cada combinación educativa observada.
            conteos_condact[condicion] += 1  # Cuenta los códigos laborales para auditar su dominio.
            if nivel in {'8', '9', '10'} and titulo not in {'1', '2'}:  # Detecta códigos no válidos de título en educación superior.
                faltantes_titulo_mes[mes] += 1  # Registra los casos que requieren revisión.
            try:  # Intenta interpretar la edad como número entero.
                edad = int(fila['p03'].strip())  # Convierte el texto de edad en un número.
            except ValueError:  # Trata una edad ilegible como desconocida, nunca como cero.
                edades_invalidas += 1  # Cuenta el problema de calidad.
                continue  # No decide elegibilidad para una edad desconocida.
            if edad >= 15 and nivel in {'8', '9', '10'} and titulo == '1' and condicion in {'1', '2', '3', '4', '5', '6', '7', '8'}:  # Aplica la regla preliminar documentada.
                candidatos_mes[mes] += 1  # Cuenta registros candidatos, no personas únicas ni población expandida.

# BLOQUE 4: registrar trazabilidad y resultados agregados de la comprobación.
huella = hashlib.sha256(ARCHIVO_ZIP.read_bytes()).hexdigest()  # Calcula la firma SHA-256 del ZIP original.
reporte = {  # Reúne evidencia técnica sin exportar registros individuales.
    'ejecutado_utc': datetime.now(timezone.utc).isoformat(),  # Registra el momento real de ejecución.
    'fuente': 'https://www.ecuadorencifras.gob.ec/documentos/web-inec/EMPLEO/2025/anual/2_BDD_DATOS_ABIERTOS_ENEMDU_2025_CSV.zip',  # Guarda la procedencia.
    'sha256_zip': huella,  # Guarda la firma para comprobar futuras copias.
    'integridad_crc': 'correcta',  # Informa que todos los miembros superaron la revisión.
    'inventario': inventario,  # Incluye los nombres y tamaños originales.
    'numero_columnas': len(columnas),  # Informa cuántas variables contiene el CSV de personas.
    'columnas_sin_documentar': sorted(set(columnas) - set(variables_documentadas)),  # Compara datos con diccionario.
    'variables_documentadas_ausentes': sorted(set(variables_documentadas) - set(columnas)),  # Revisa la diferencia inversa.
    'registros_personas': sum(conteos_mes.values()),  # Informa filas, sin asumir personas distintas.
    'registros_por_mes': dict(sorted(conteos_mes.items())),  # Permite revisar la cobertura anual.
    'nivel_por_titulo': [{'p10a': clave[0], 'p12a': clave[1], 'n': valor} for clave, valor in sorted(conteos_educacion.items())],  # Documenta consistencia educativa.
    'condact_codigos': dict(sorted(conteos_condact.items())),  # Documenta el dominio laboral observado.
    'candidatos_por_mes': dict(sorted(candidatos_mes.items())),  # Informa elegibilidad preliminar sin ponderación.
    'candidatos_total': sum(candidatos_mes.values()),  # Resume registros candidatos al estudio.
    'titulo_invalido_o_faltante_en_superior_por_mes': dict(sorted(faltantes_titulo_mes.items())),  # Informa la calidad de p12a.
    'edades_no_interpretables': edades_invalidas,  # Informa problemas de conversión de edad.
    'limite': 'Auditoría preliminar; no verifica personas únicas, diseño completo ni asociaciones. No es una estimación poblacional.',  # Delimita el alcance.
}  # Cierra el conjunto de resultados.
ruta_reporte = SALIDAS / '00_verificacion_fuentes.json'  # Define el archivo reproducible de salida.
ruta_reporte.write_text(json.dumps(reporte, ensure_ascii=False, indent=2), encoding='utf-8')  # Guarda el reporte legible.
print(f'Registros originales: {reporte["registros_personas"]:,}')  # Muestra el tamaño observado.
print(f'Columnas: {reporte["numero_columnas"]}')  # Muestra el número de variables.
print(f'Registros candidatos: {reporte["candidatos_total"]:,}')  # Muestra el tamaño preliminar del dominio.
print(f'Reporte guardado: {ruta_reporte}')  # Indica dónde consultar la evidencia completa.
